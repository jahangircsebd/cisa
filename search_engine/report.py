"""Word (.docx) report generation."""
from __future__ import annotations

from collections import Counter
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .engine import SearchRun
from .models import Result

CATEGORY_ORDER = ["News", "Social", "Media", "Academic", "Web"]


def _hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _table(doc, header: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    return table


def _result_entry(doc, i: int, r: Result) -> None:
    p = doc.add_paragraph(style="List Number")
    title = p.add_run(r.title or r.url)
    title.bold = True
    meta = [r.platform or r.domain, r.published or "date unknown",
            f"relevance {r.relevance:.2f}"]
    if r.location_hits:
        meta.append("location: " + ", ".join(r.location_hits))
    if r.name_on_page:
        meta.append("name confirmed on page")
    elif r.page_status and not r.page_status.startswith("ok"):
        meta.append("page not readable")
    m = doc.add_paragraph()
    mr = m.add_run(" | ".join(meta))
    mr.italic = True
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    if r.snippet:
        s = doc.add_paragraph(r.snippet[:500])
        s.paragraph_format.left_indent = Pt(18)
    if r.page_excerpt and r.page_excerpt[:300] != r.snippet[:300]:
        e = doc.add_paragraph()
        e.paragraph_format.left_indent = Pt(18)
        lead = e.add_run("On the page: ")
        lead.bold = True
        lead.font.size = Pt(9)
        body = e.add_run(r.page_excerpt[:600])
        body.font.size = Pt(9)
    link = doc.add_paragraph()
    link.paragraph_format.left_indent = Pt(18)
    _hyperlink(link, r.url, r.url)


def build_report(run: SearchRun, path: str, analyst_notes: str | None = None) -> str:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Keyword Media Search Report", 0)
    _table(doc, ["Field", "Value"], [
        ["Keyword / phrase", run.keyword],
        ["Name variants searched", ", ".join(run.aliases) or "-"],
        ["Location focus", run.location or "None (global)"],
        ["Search run (UTC)", run.started],
        ["Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
        ["Queries issued", "; ".join(run.queries) or "-"],
        ["Total unique results", str(len(run.results))],
    ])

    # Executive summary
    doc.add_heading("1. Executive Summary", 1)
    local = [r for r in run.results if r.location_score > 0]
    cats = Counter(r.category for r in run.results)
    plats = Counter(r.platform for r in run.results if r.platform)
    doc.add_paragraph(
        f"The search for “{run.keyword}” returned {len(run.results)} unique, "
        f"keyword-relevant items across {len(set(r.domain for r in run.results))} domains. "
        + (f"{len(local)} of them carry explicit {run.location} location signals "
           f"(state name, city, institution or local outlet)." if run.location else "")
    )
    if analyst_notes:
        doc.add_paragraph(analyst_notes)
    if run.results:
        top = run.results[0]
        doc.add_paragraph(f"Top-ranked item: {top.title} ({top.domain}).", style="List Bullet")
    for c in CATEGORY_ORDER:
        if cats.get(c):
            doc.add_paragraph(f"{c}: {cats[c]} item(s)", style="List Bullet")
    if plats:
        doc.add_paragraph("Social platforms: " + ", ".join(f"{k} ({v})" for k, v in plats.most_common()),
                          style="List Bullet")

    # Breakdown tables
    doc.add_heading("2. Coverage Breakdown", 1)
    doc.add_heading("By category", 2)
    _table(doc, ["Category", "Items", f"With {run.location or 'location'} signal"],
           [[c, cats.get(c, 0), sum(1 for r in local if r.category == c)] for c in CATEGORY_ORDER])
    doc.add_heading("Top domains", 2)
    _table(doc, ["Domain", "Items"], [[d, n] for d, n in
                                      Counter(r.domain for r in run.results).most_common(15)])
    if run.location:
        hits = Counter(h for r in run.results for h in r.location_hits)
        if hits:
            doc.add_heading(f"{run.location} location signals found", 2)
            _table(doc, ["Signal", "Occurrences"], [[h, n] for h, n in hits.most_common(20)])

    # Detailed results
    doc.add_heading("3. Detailed Results", 1)
    sec = 1
    groups = [(f"{run.location}-linked", local), ("Other relevant", [r for r in run.results if r.location_score == 0])] \
        if run.location else [("All", run.results)]
    for label, items in groups:
        for cat in CATEGORY_ORDER:
            subset = [r for r in items if r.category == cat]
            if not subset:
                continue
            doc.add_heading(f"3.{sec} {label} – {cat} ({len(subset)})", 2)
            sec += 1
            for i, r in enumerate(subset, 1):
                _result_entry(doc, i, r)

    # Methodology
    doc.add_heading("4. Sources & Methodology", 1)
    _table(doc, ["Source", "Status"], [[k, v] for k, v in sorted(run.source_status.items())])
    doc.add_paragraph()
    for line in [
        "Each source was queried with the keyword and each name variant, with and without the location term. "
        "Academic databases (OpenAlex, Crossref, Semantic Scholar) were searched by author name.",
        "Local sources searched only inside the location's own news outlets and institutions.",
        "The top results' pages were opened and read in full (including PDFs). Text around name, topic and "
        "location mentions was used for scoring; pages that were readable but never mention the name were dropped.",
        "Results were de-duplicated by normalized URL (tracking parameters removed).",
        "Keyword score = share of keyword terms present in title/snippet/URL/page text (+0.2 for exact phrase); "
        "a name variant plus the topic words counts as a full match.",
        "Location score reflects mentions of the state, its cities/counties/institutions, or local outlet domains.",
        "Relevance = 0.6 × keyword score + 0.4 × location score.",
        "Social media coverage (TikTok, Instagram, Facebook, X, LinkedIn) is limited to publicly indexed "
        "posts/profiles; name matches on social platforms may refer to different people and should be verified.",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.save(path)
    return path
